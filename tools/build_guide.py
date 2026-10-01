"""Genera el taller guiado a partir de data/taller_guiado.json.

Salidas:
  assets/js/guide.js                      -> datos para la sección «Taller guiado» del sitio
  assets/docs/InfraLab_Taller_Guiado.docx -> versión editable (Word) para que los equipos la completen
  assets/docs/InfraLab_Taller_Guiado.pdf  -> versión para imprimir (se genera con Chrome/Edge headless)

Uso:  python tools/build_guide.py
Requiere: python-docx  (pip install python-docx)
"""
import json
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "taller_guiado.json"
JS_OUT = ROOT / "assets" / "js" / "guide.js"
DOCS = ROOT / "assets" / "docs"
DOCX_OUT = DOCS / "InfraLab_Taller_Guiado.docx"
PDF_OUT = DOCS / "InfraLab_Taller_Guiado.pdf"

ACCENT = RGBColor(0x0E, 0x74, 0x90)
MUTED = RGBColor(0x55, 0x60, 0x70)
RED = RGBColor(0xB4, 0x23, 0x18)


# ------------------------------------------------------------------ utilidades docx
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=10, color=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def bottom_border(paragraph, color="0E7490", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), size)
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    bdr.append(b)
    p_pr.append(bdr)


def tour_url(g, tid):
    base = (g.get("site_url") or "").strip()
    if base and not base.endswith("/"):
        base += "/"
    return f"{base}index.html#guia/{tid}" if base else f"index.html#guia/{tid}"


def add_link(paragraph, url, text):
    """Hipervínculo real en Word (python-docx no trae un método directo)."""
    r_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0E7490"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); run.append(t)
    link.append(run)
    paragraph._p.append(link)


def label(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = ACCENT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    return p


def boxed(doc, lines, fill="FFF4E5", title=None, title_color=RED):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, fill)
    cell.text = ""
    first = True
    if title:
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = title_color
        first = False
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        r = p.add_run(line)
        r.font.size = Pt(10.5)
        r.italic = True
    doc.add_paragraph()


def worksheet(doc, cols, rows):
    t = doc.add_table(rows=rows + 1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cols):
        cell = t.rows[0].cells[i]
        shade(cell, "0E7490")
        set_cell_text(cell, c, bold=True, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r in range(1, rows + 1):
        t.rows[r].height = Cm(1.3)
    return t


# ------------------------------------------------------------------ documento
def build_docx(g):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # Carta
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin = sec.bottom_margin = Cm(2)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    footer = sec.footer.paragraphs[0]
    footer.text = "InfraLab · Taller guiado de Gestión de Infraestructura TI · Uniminuto 2026-2"
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = MUTED

    # Portada
    doc.add_paragraph().paragraph_format.space_before = Pt(80)
    p = doc.add_paragraph()
    r = p.add_run("InfraLab")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = ACCENT
    h = doc.add_heading(g["title"], level=0)
    h.runs[0].font.size = Pt(30)
    p = doc.add_paragraph(g["subtitle"])
    p.runs[0].font.size = Pt(13)
    p.runs[0].font.color.rgb = MUTED
    bottom_border(p)
    doc.add_paragraph()
    t = doc.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    for i, lab in enumerate(["Integrantes del equipo", "Grupo / NRC", "Caso asignado", "Fecha"]):
        shade(t.rows[i].cells[0], "E6F4F8")
        set_cell_text(t.rows[i].cells[0], lab, bold=True)
        t.rows[i].height = Cm(1.1)
    doc.add_paragraph()
    doc.add_heading("Cómo usar este taller", level=1)
    for line in g["intro"]:
        doc.add_paragraph(line)
    for i, s in enumerate(g["how"], 1):
        doc.add_paragraph(f"{i}. {s}")

    # Mapa de etapas
    doc.add_heading("Las diez etapas", level=1)
    t = doc.add_table(rows=len(g["steps"]) + 1, cols=3)
    t.style = "Table Grid"
    for i, c in enumerate(["Etapa", "Para qué sirve", "Caso de ejemplo"]):
        shade(t.rows[0].cells[i], "0E7490")
        set_cell_text(t.rows[0].cells[i], c, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r_i, s in enumerate(g["steps"], 1):
        set_cell_text(t.rows[r_i].cells[0], f"{s['id']}. {s['title']}", bold=True, size=9.5)
        set_cell_text(t.rows[r_i].cells[1], s["discover"], size=9.5)
        set_cell_text(t.rows[r_i].cells[2], s["case_label"], size=9.5)

    # Herramientas del sitio
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Herramientas del sitio que usarás", level=1)
    p = doc.add_paragraph("Todas están en el sitio InfraLab. En cada hoja «Tu turno» se indica cuál usar, qué hacer y qué trampa evitar.")
    p.runs[0].font.color.rgb = MUTED
    tk = g.get("toolkit", [])
    t = doc.add_table(rows=len(tk) + 1, cols=4)
    t.style = "Table Grid"
    for i, c in enumerate(["Herramienta", "Qué es", "Para qué la usas", "Etapas"]):
        shade(t.rows[0].cells[i], "0E7490")
        set_cell_text(t.rows[0].cells[i], c, bold=True, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r_i, k in enumerate(tk, 1):
        set_cell_text(t.rows[r_i].cells[0], k["name"], bold=True, size=9)
        set_cell_text(t.rows[r_i].cells[1], k["what"], size=9)
        set_cell_text(t.rows[r_i].cells[2], k["use"], size=9)
        set_cell_text(t.rows[r_i].cells[3], "Todas" if len(k["steps"]) == len(g["steps"]) else ", ".join(map(str, k["steps"])), size=9)
    TK = {k["id"]: k for k in tk}

    # Ejercicios guiados en la web
    tours = g.get("tours", [])
    TR = {t["id"]: t for t in tours}
    if tours:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Paso 1 · Ejercicios guiados en la web", level=1)
        p = doc.add_paragraph("Hazlos antes de las etapas. Abre el enlace: el sitio resalta dónde hacer clic, comprueba cada acción y al final te explica qué encontraste. Luego repite lo mismo con tu caso.")
        p.runs[0].font.color.rgb = MUTED
        if not (g.get("site_url") or "").strip():
            p = doc.add_paragraph("Si el enlace no abre, entra al sitio, ve a «Taller guiado» y elige el ejercicio por su nombre.")
            p.runs[0].font.size = Pt(9); p.runs[0].italic = True
        for k, t in enumerate(tours, 1):
            h = doc.add_heading(f"Ejercicio {k} · {t['title']} ({t['minutes']} min · {t.get('group', '')})", level=2)
            p = doc.add_paragraph()
            r = p.add_run("Enlace: "); r.bold = True
            add_link(p, tour_url(g, t["id"]), tour_url(g, t["id"]))
            p = doc.add_paragraph()
            r = p.add_run("Para qué: "); r.bold = True
            p.add_run(t["purpose"])
            for i, s in enumerate(t["steps"], 1):
                q = doc.add_paragraph(f"{i}. {s['t']}")
                q.paragraph_format.left_indent = Cm(0.6)
                q.runs[0].font.size = Pt(9.5)
                q.paragraph_format.space_after = Pt(1)
            p = doc.add_paragraph()
            r = p.add_run("Tu turno: "); r.bold = True; r.font.color.rgb = ACCENT
            p.add_run(t["yours"])

    # Etapas
    for s in g["steps"]:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = doc.add_paragraph()
        r = p.add_run(f"ETAPA {s['id']} DE {len(g['steps'])}")
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = ACCENT
        doc.add_heading(s["title"], level=1)
        p = doc.add_paragraph()
        r = p.add_run(f"Caso de ejemplo: {s['case_label']}   ·   En el sitio: {s['tab']}")
        r.font.size = Pt(9.5)
        r.font.color.rgb = MUTED

        if s.get("tours"):
            p = doc.add_paragraph()
            r = p.add_run("Empieza por los ejercicios guiados en la web: "); r.bold = True; r.font.color.rgb = ACCENT
            for j, tid in enumerate(s["tours"]):
                if j: p.add_run("  ·  ")
                add_link(p, tour_url(g, tid), TR[tid]["title"])
        label(doc, "Para qué sirve")
        doc.add_paragraph(s["purpose"])
        p = doc.add_paragraph()
        r = p.add_run("Lo que vas a descubrir: ")
        r.bold = True
        p.add_run(s["discover"])

        label(doc, "Qué haces")
        for a in s["activities"]:
            doc.add_paragraph(a, style="List Bullet")

        label(doc, f"Detecta los errores · {s['case_label']}")
        boxed(doc, s["attempt"], title=f"Intento de un equipo (contiene {len(s['errors'])} errores). Encuéntralos antes de leer la clave del anexo.")
        doc.add_paragraph("Errores que encontramos:")
        for _ in range(len(s["errors"])):
            p = doc.add_paragraph("______________________________________________________________________________")
            p.runs[0].font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)

        label(doc, "Autoverificación")
        for c in s["checks"]:
            doc.add_paragraph(f"☐  {c}")

        # Hoja de trabajo en página propia
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = doc.add_paragraph()
        r = p.add_run(f"ETAPA {s['id']} · HOJA DE TRABAJO")
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = ACCENT
        doc.add_heading(f"Tu turno: {s['title']}", level=1)
        p = doc.add_paragraph("Aplica esta etapa a tu caso. Recuerda: problema → evidencia → impacto → decisión → métrica.")
        p.runs[0].font.color.rgb = MUTED
        label(doc, "Herramientas del sitio para esta etapa")
        tt = doc.add_table(rows=len(s.get("tools", [])) + 1, cols=3)
        tt.style = "Table Grid"
        for i, c in enumerate(["Herramienta", "Qué haces", "Trampa a evitar"]):
            shade(tt.rows[0].cells[i], "155E75")
            set_cell_text(tt.rows[0].cells[i], c, bold=True, size=9, color=RGBColor(0xFF, 0xFF, 0xFF))
        for r_i, x in enumerate(s.get("tools", []), 1):
            set_cell_text(tt.rows[r_i].cells[0], TK.get(x["tool"], {}).get("name", x["tool"]), bold=True, size=8.5)
            set_cell_text(tt.rows[r_i].cells[1], x["do"], size=8.5)
            set_cell_text(tt.rows[r_i].cells[2], x["trap"], size=8.5, color=RED)
        label(doc, "Tu tabla de trabajo")
        worksheet(doc, s["worksheet"]["cols"], s["worksheet"]["rows"])
        label(doc, "Evidencia del caso que usamos (dato, página o pestaña del sitio)")
        boxed(doc, [" ", " "], fill="FFFFFF")
        label(doc, "¿Para qué nos sirvió esta etapa? ¿Qué cambia en las siguientes?")
        boxed(doc, [" ", " "], fill="FFFFFF")

    # Anexo: clave de errores
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Anexo · Clave de errores", level=1)
    doc.add_paragraph("Léela después de intentar cada ejercicio. La clave explica qué está mal y qué revisar; no entrega la solución del caso.").runs[0].italic = True
    for s in g["steps"]:
        doc.add_heading(f"Etapa {s['id']} · {s['title']} ({s['case_label']})", level=2)
        for i, e in enumerate(s["errors"], 1):
            p = doc.add_paragraph()
            r = p.add_run(f"{i}. {e['e']} ")
            r.bold = True
            p.add_run(e["why"])

    # La plantilla base de python-docx trae <w:zoom> sin el atributo obligatorio «percent».
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    DOCS.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


def find_browser():
    candidates = [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        "google-chrome", "chromium", "chromium-browser", "microsoft-edge",
    ]
    for c in candidates:
        if Path(c).exists() or shutil.which(c):
            return c
    return None


def build_pdf():
    browser = find_browser()
    page = ROOT / "taller-guiado.html"
    if not browser or not page.exists():
        print("  (PDF omitido: no se encontró Chrome/Edge o falta taller-guiado.html)")
        return
    subprocess.run([browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--virtual-time-budget=8000",
                    f"--print-to-pdf={PDF_OUT}", page.as_uri() + "?print=1"], check=False, capture_output=True)


def main():
    g = json.loads(SRC.read_text(encoding="utf-8"))
    JS_OUT.write_text("/* Archivo generado por tools/build_guide.py — no editar a mano. */\nwindow.GUIDE = "
                      + json.dumps(g, ensure_ascii=False, indent=1) + ";\n", encoding="utf-8")
    print(f"OK -> {JS_OUT.relative_to(ROOT)}")
    build_docx(g)
    print(f"OK -> {DOCX_OUT.relative_to(ROOT)}")
    build_pdf()
    if PDF_OUT.exists():
        print(f"OK -> {PDF_OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
